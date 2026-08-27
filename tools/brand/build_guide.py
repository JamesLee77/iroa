from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from pathlib import Path
import re
from tempfile import NamedTemporaryFile
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
BRAND_ROOT = REPOSITORY_ROOT / "docs/brand"
DEFAULT_SOURCE = BRAND_ROOT / "IROA_BI_GUIDE_KO.md"
DEFAULT_OUTPUT = BRAND_ROOT / "IROA_BI_GUIDE_KO.docx"
FONT_ROOT = BRAND_ROOT / "assets/fonts"

TITLE = "IROA.AI Brand Identity Guide"
SUBJECT = "IROA.AI 브랜드 아이덴티티 가이드 · v1.0 · 2026-08-27"
VERSION = "1.0"
REFERENCE_DATE = "2026-08-27"
STYLE_CATEGORY = "compact_reference_guide / editorial_cover / A4 / IROA palette"

# compact_reference_guide resolved to A4 with named IROA overrides.
PAGE_WIDTH_DXA = 11906
PAGE_HEIGHT_DXA = 16838
MARGIN_TOP_DXA = 1134
MARGIN_SIDE_DXA = 1020
MARGIN_BOTTOM_DXA = 1020
HEADER_FOOTER_DXA = 567
CONTENT_WIDTH_DXA = 9866
TABLE_INDENT_DXA = 120
CELL_MARGIN_VERTICAL_DXA = 80
CELL_MARGIN_HORIZONTAL_DXA = 120

FONT_NAME = "Noto Sans KR"
FONT_HASHES = {
    "NotoSansKR-Medium.otf": "b46988ef13e8bac08f3933af686eaf770972994f9b6d335be0184d60169b5431",
    "NotoSansKR-Bold.otf": "5a6ceb287ed2fc6cfc6213144ebea68cbd94b20fc9eb873d8486493bf02d9bda",
}

NAVY = "16263D"
CORAL = "F06D5E"
IVORY = "F7F3EA"
TEAL = "3D8B83"
LIGHT_TEAL = "83CDC4"
INK = "19222E"
WHITE = "FFFFFF"

PAGE_BREAK_HEADINGS: set[str] = set()
DOCX_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


def _append(parent, tag: str, **attributes: str):
    node = OxmlElement(tag)
    for name, value in attributes.items():
        node.set(qn(name), str(value))
    parent.append(node)
    return node


def _set_font(run, *, size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for role in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{role}"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _set_style(
    style,
    *,
    size: float,
    color: str,
    before: float,
    after: float,
    line: float,
    bold: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
) -> None:
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    properties = style.element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for role in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{role}"), None)
    for role in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{role}"), FONT_NAME)
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line
    paragraph.widow_control = True
    if alignment is not None:
        paragraph.alignment = alignment


def _configure_styles(document: Document) -> None:
    styles = document.styles
    _set_style(styles["Normal"], size=10.5, color=INK, before=0, after=6, line=1.25)
    _set_style(
        styles["Title"],
        size=28,
        color=NAVY,
        before=0,
        after=12,
        line=1.25,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_style(styles["Subtitle"], size=13, color=NAVY, before=0, after=10, line=1.25)
    _set_style(styles["Heading 1"], size=16, color=NAVY, before=18, after=10, line=1.25, bold=True)
    _set_style(styles["Heading 2"], size=13, color=NAVY, before=14, after=7, line=1.25, bold=True)
    _set_style(styles["Heading 3"], size=12, color=INK, before=10, after=5, line=1.25, bold=True)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        styles[name].paragraph_format.keep_with_next = True
        styles[name].paragraph_format.keep_together = True

    custom = {
        "IROA Kicker": (9.5, NAVY, 0, 10, 1.15, True, WD_ALIGN_PARAGRAPH.CENTER),
        "IROA Caption": (8.5, INK, 3, 9, 1.15, False, WD_ALIGN_PARAGRAPH.CENTER),
        "IROA Note": (10, INK, 5, 7, 1.25, False, None),
        "IROA List Bullet": (10.5, INK, 0, 4, 1.25, False, None),
        "IROA List Number": (10.5, INK, 0, 4, 1.25, False, None),
        "IROA Table Header": (9, WHITE, 0, 0, 1.15, True, None),
        "IROA Table Body": (9, INK, 0, 0, 1.15, False, None),
    }
    for name, tokens in custom.items():
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_style(
            style,
            size=tokens[0],
            color=tokens[1],
            before=tokens[2],
            after=tokens[3],
            line=tokens[4],
            bold=tokens[5],
            alignment=tokens[6],
        )

    table_style = styles["IROA Table"] if "IROA Table" in styles else styles.add_style("IROA Table", WD_STYLE_TYPE.TABLE)
    table_style.font.name = FONT_NAME
    table_style.font.size = Pt(9)
    table_style.font.color.rgb = RGBColor.from_string(INK)
    table_fonts = table_style.element.get_or_add_rPr().get_or_add_rFonts()
    for role in ("ascii", "hAnsi", "eastAsia", "cs"):
        table_fonts.set(qn(f"w:{role}"), FONT_NAME)


def _configure_section(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Twips(PAGE_WIDTH_DXA)
    section.page_height = Twips(PAGE_HEIGHT_DXA)
    section.top_margin = Twips(MARGIN_TOP_DXA)
    section.right_margin = Twips(MARGIN_SIDE_DXA)
    section.bottom_margin = Twips(MARGIN_BOTTOM_DXA)
    section.left_margin = Twips(MARGIN_SIDE_DXA)
    section.header_distance = Twips(HEADER_FOOTER_DXA)
    section.footer_distance = Twips(HEADER_FOOTER_DXA)
    section.different_first_page_header_footer = True


def _add_field(paragraph, instruction: str, placeholder: str) -> None:
    run = paragraph.add_run()
    _set_font(run, size=8.5, color=NAVY)
    _append(run._r, "w:fldChar", **{"w:fldCharType": "begin"})
    instruction_run = paragraph.add_run()
    _set_font(instruction_run, size=8.5, color=NAVY)
    text = _append(instruction_run._r, "w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separator = paragraph.add_run()
    _set_font(separator, size=8.5, color=NAVY)
    _append(separator._r, "w:fldChar", **{"w:fldCharType": "separate"})
    shown = paragraph.add_run(placeholder)
    _set_font(shown, size=8.5, color=NAVY)
    end = paragraph.add_run()
    _set_font(end, size=8.5, color=NAVY)
    _append(end._r, "w:fldChar", **{"w:fldCharType": "end"})


def _configure_headers_and_footers(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("IROA.AI 브랜드 아이덴티티 가이드 · V1.0")
    _set_font(run, size=8.5, color=NAVY, bold=True)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    prefix = paragraph.add_run(f"{REFERENCE_DATE}  |  페이지 ")
    _set_font(prefix, size=8.5, color=NAVY)
    _add_field(paragraph, "PAGE", "2")
    slash = paragraph.add_run(" / ")
    _set_font(slash, size=8.5, color=NAVY)
    _add_field(paragraph, "NUMPAGES", "1")

    first_footer = section.first_page_footer
    first_paragraph = first_footer.paragraphs[0]
    first_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = first_paragraph.add_run("IROA.AI · 1.0 / 2026-08-27")
    _set_font(run, size=8.5, color=NAVY, bold=True)


def _add_hyperlink(paragraph, label: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for role in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{role}"), FONT_NAME)
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), NAVY)
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _add_inline(paragraph, source: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(source):
        if match.start() > cursor:
            paragraph.add_run(source[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_font(run, color=NAVY, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            _set_font(run, color=NAVY)
        else:
            label, target = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            _add_hyperlink(paragraph, label, target)
        cursor = match.end()
    if cursor < len(source):
        paragraph.add_run(source[cursor:])


def _add_paragraph(document: Document, text: str, style: str = "Normal"):
    paragraph = document.add_paragraph(style=style)
    _add_inline(paragraph, text)
    return paragraph


def _add_rule(paragraph, *, color: str = CORAL, size: str = "18") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _add_image(document: Document, path: Path, *, width: float, alt_text: str, caption: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = caption is not None
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", Path(alt_text.split(" · ", 1)[0]).name)
    if caption:
        _add_paragraph(document, caption, "IROA Caption")


def _list_numbering(document: Document) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_base = max(abstract_ids, default=0) + 1
    num_base = max(num_ids, default=0) + 1

    def add_abstract(identifier: int, kind: str, marker: str) -> None:
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(identifier))
        _append(abstract, "w:multiLevelType", **{"w:val": "singleLevel"})
        level = _append(abstract, "w:lvl", **{"w:ilvl": "0"})
        _append(level, "w:start", **{"w:val": "1"})
        _append(level, "w:numFmt", **{"w:val": kind})
        _append(level, "w:lvlText", **{"w:val": marker})
        _append(level, "w:lvlJc", **{"w:val": "left"})
        paragraph_properties = _append(level, "w:pPr")
        tabs = _append(paragraph_properties, "w:tabs")
        _append(tabs, "w:tab", **{"w:val": "num", "w:pos": "269"})
        _append(paragraph_properties, "w:ind", **{"w:left": "540", "w:hanging": "271"})
        spacing = _append(paragraph_properties, "w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        run_properties = _append(level, "w:rPr")
        _append(
            run_properties,
            "w:rFonts",
            **{
                "w:ascii": FONT_NAME,
                "w:hAnsi": FONT_NAME,
                "w:eastAsia": FONT_NAME,
                "w:cs": FONT_NAME,
            },
        )
        _append(run_properties, "w:color", **{"w:val": NAVY})
        _append(run_properties, "w:sz", **{"w:val": "21"})
        numbering.append(abstract)

    def add_num(identifier: int, abstract_identifier: int) -> None:
        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(identifier))
        _append(number, "w:abstractNumId", **{"w:val": str(abstract_identifier)})
        numbering.append(number)

    add_abstract(abstract_base, "bullet", "•")
    add_abstract(abstract_base + 1, "decimal", "%1.")
    add_num(num_base, abstract_base)
    add_num(num_base + 1, abstract_base + 1)
    return num_base, num_base + 1


def _apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    number_properties = properties.find(qn("w:numPr"))
    if number_properties is None:
        number_properties = OxmlElement("w:numPr")
        properties.append(number_properties)
    _append(number_properties, "w:ilvl", **{"w:val": "0"})
    _append(number_properties, "w:numId", **{"w:val": str(num_id)})


def _visible_width(value: str) -> int:
    plain = re.sub(r"\*\*|`", "", value)
    plain = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", plain)
    return sum(2 if ord(character) > 127 else 1 for character in plain)


def _column_widths(rows: list[list[str]]) -> list[int]:
    count = len(rows[0])
    scores = []
    for index in range(count):
        score = max(_visible_width(row[index]) for row in rows)
        scores.append(max(6, min(score, 42)))
    minimum = 760 if count >= 4 else 980
    remaining = CONTENT_WIDTH_DXA - minimum * count
    total = sum(scores)
    widths = [minimum + round(remaining * score / total) for score in scores]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def _replace_child(parent, tag: str, replacement) -> None:
    for child in list(parent):
        if child.tag == qn(tag):
            parent.remove(child)
    parent.append(replacement)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:shd"))
    if existing is not None:
        properties.remove(existing)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _configure_table(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "IROA Table"
    properties = table._tbl.tblPr

    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    width.set(qn("w:type"), "dxa")
    _replace_child(properties, "w:tblW", width)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    _replace_child(properties, "w:tblInd", indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    _replace_child(properties, "w:tblLayout", layout)

    margins = OxmlElement("w:tblCellMar")
    for name, value in (
        ("top", CELL_MARGIN_VERTICAL_DXA),
        ("bottom", CELL_MARGIN_VERTICAL_DXA),
        ("start", CELL_MARGIN_HORIZONTAL_DXA),
        ("end", CELL_MARGIN_HORIZONTAL_DXA),
    ):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    _replace_child(properties, "w:tblCellMar", margins)

    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), NAVY)
        borders.append(node)
    _replace_child(properties, "w:tblBorders", borders)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    _append(header_properties, "w:tblHeader", **{"w:val": "true"})
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        _append(row_properties, "w:cantSplit", **{"w:val": "true"})
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width = widths[column_index]
            cell.width = Twips(cell_width)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width_node = OxmlElement("w:tcW")
            cell_width_node.set(qn("w:w"), str(cell_width))
            cell_width_node.set(qn("w:type"), "dxa")
            _replace_child(cell_properties, "w:tcW", cell_width_node)
            _shade(cell, NAVY if row_index == 0 else (IVORY if row_index % 2 == 0 else WHITE))


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    widths = _column_widths(rows)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.style = "IROA Table Header" if row_index == 0 else "IROA Table Body"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _add_inline(paragraph, value)
            if row_index == 0:
                for run in paragraph.runs:
                    _set_font(run, size=9, color=WHITE, bold=True)
    _configure_table(table, widths)


def _parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    if rows and any(len(row) != len(rows[0]) for row in rows):
        raise ValueError(f"inconsistent Markdown table at line {start + 1}")
    return rows, index


def _add_cover(document: Document, lines: list[str], bullet_num_id: int) -> int:
    if not lines or not lines[0].startswith("# "):
        raise ValueError("guide must start with one Markdown title")
    _add_image(
        document,
        BRAND_ROOT / "exports/digital/iroa-wordmark-512.png",
        width=3.25,
        alt_text="exports/digital/iroa-wordmark-512.png · Track A 공식 컬러 워드마크",
    )
    kicker = _add_paragraph(document, "단일 콘텐츠 원본 · Track A", "IROA Kicker")
    kicker.paragraph_format.space_before = Pt(22)
    title = _add_paragraph(document, lines[0][2:].strip(), "Title")
    _add_rule(title, color=CORAL, size="18")

    index = 1
    while index < len(lines) and not lines[index].strip():
        index += 1
    while index < len(lines) and lines[index].startswith("- "):
        paragraph = _add_paragraph(document, lines[index][2:].strip(), "IROA List Bullet")
        _apply_numbering(paragraph, bullet_num_id)
        index += 1
    while index < len(lines) and not lines[index].strip():
        index += 1
    if index < len(lines) and not lines[index].startswith("#"):
        paragraph = _add_paragraph(document, lines[index])
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(0)
        index += 1
    document.add_page_break()
    return index


def _add_section_image(document: Document, heading: str) -> None:
    if heading == "1. 브랜드 핵심":
        _add_image(
            document,
            BRAND_ROOT / "examples/usage-overview.png",
            width=6.70,
            alt_text="examples/usage-overview.png · Track A 공식 자산 적용 예시 보드",
            caption="공식 사용 예시 보드 · examples/usage-overview.png",
        )
    elif heading == "2. 공식 로고 시스템":
        _add_image(
            document,
            BRAND_ROOT / "exports/digital/iroa-symbol-512.png",
            width=1.30,
            alt_text="exports/digital/iroa-symbol-512.png · Track A 공식 컬러 심볼",
            caption="공식 컬러 심볼 · exports/digital/iroa-symbol-512.png",
        )
    elif heading == "6. 디지털 아이콘":
        _add_image(
            document,
            BRAND_ROOT / "exports/icons/app-icon-192.png",
            width=1.15,
            alt_text="exports/icons/app-icon-192.png · IROA 앱 아이콘",
            caption="앱 아이콘 · exports/icons/app-icon-192.png",
        )
    elif heading == "9. 매체별 적용":
        _add_image(
            document,
            BRAND_ROOT / "exports/icons/symbol-kiosk-1024.png",
            width=1.30,
            alt_text="exports/icons/symbol-kiosk-1024.png · IROA 키오스크 심볼",
            caption="키오스크 심볼 · exports/icons/symbol-kiosk-1024.png",
        )


def _add_markdown_body(document: Document, lines: list[str], start: int, bullet_num_id: int, decimal_num_id: int) -> None:
    index = start
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows, index = _parse_table(lines, index)
            _add_table(document, rows)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1)) - 1
            text = heading.group(2)
            paragraph = _add_paragraph(document, text, f"Heading {level}")
            if text in PAGE_BREAK_HEADINGS:
                paragraph.paragraph_format.page_break_before = True
            _add_section_image(document, text)
            index += 1
            continue
        if line.startswith("> "):
            paragraph = _add_paragraph(document, line[2:], "IROA Note")
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:fill"), IVORY)
            properties.append(shading)
            borders = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), CORAL)
            borders.append(left)
            properties.append(borders)
            index += 1
            continue
        if line.startswith("- "):
            paragraph = _add_paragraph(document, line[2:], "IROA List Bullet")
            _apply_numbering(paragraph, bullet_num_id)
            index += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            paragraph = _add_paragraph(document, numbered.group(1), "IROA List Number")
            _apply_numbering(paragraph, decimal_num_id)
            index += 1
            continue
        _add_paragraph(document, line)
        index += 1


def _enable_field_updates(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _validate_fonts() -> None:
    for filename, expected in FONT_HASHES.items():
        path = FONT_ROOT / filename
        if not path.is_file():
            raise FileNotFoundError(f"pinned font missing: {path}")
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            raise ValueError(f"pinned font hash mismatch: {path}: {actual}")


def _obfuscate_font(payload: bytes, font_key: str) -> bytes:
    key = bytes.fromhex(font_key.strip("{}").replace("-", ""))[::-1]
    obfuscated = bytearray(payload)
    for index in range(min(32, len(obfuscated))):
        obfuscated[index] ^= key[index % 16]
    return bytes(obfuscated)


def _embed_pinned_fonts(path: Path) -> None:
    regular_key = "{B46988EF-13E8-BAC0-8F39-33AF686EAF77}"
    bold_key = "{5A6CEB28-7ED2-FC6C-FC62-13144EBEA68C}"
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    relationship_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
    office_relationship_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    with zipfile.ZipFile(path) as source:
        entries = {name: source.read(name) for name in source.namelist()}

    content_types = re.sub(
        rb"</Types>",
        b'<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/></Types>',
        entries["[Content_Types].xml"],
        count=1,
    )
    entries["[Content_Types].xml"] = content_types

    font_table = OxmlElement("w:fonts")
    if "word/fontTable.xml" in entries:
        font_table = etree.fromstring(entries["word/fontTable.xml"])
    noto = None
    for candidate in font_table.findall(qn("w:font")):
        if candidate.get(qn("w:name")) == FONT_NAME:
            noto = candidate
            break
    if noto is None:
        noto = OxmlElement("w:font")
        noto.set(qn("w:name"), FONT_NAME)
        font_table.append(noto)
    for child in list(noto):
        if child.tag in (qn("w:embedRegular"), qn("w:embedBold")):
            noto.remove(child)
    for role, relationship_id, key in (
        ("embedRegular", "rIdIROARegular", regular_key),
        ("embedBold", "rIdIROABold", bold_key),
    ):
        node = OxmlElement(f"w:{role}")
        node.set(qn("r:id"), relationship_id)
        node.set(qn("w:fontKey"), key)
        noto.append(node)
    entries["word/fontTable.xml"] = etree.tostring(
        font_table,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    relationships = etree.Element(f"{{{relationship_namespace}}}Relationships", nsmap={None: relationship_namespace})
    for relationship_id, target in (
        ("rIdIROARegular", "fonts/NotoSansKR-Medium.odttf"),
        ("rIdIROABold", "fonts/NotoSansKR-Bold.odttf"),
    ):
        node = etree.SubElement(relationships, f"{{{relationship_namespace}}}Relationship")
        node.set("Id", relationship_id)
        node.set("Type", relationship_type)
        node.set("Target", target)
    entries["word/_rels/fontTable.xml.rels"] = etree.tostring(
        relationships,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    entries["word/fonts/NotoSansKR-Medium.odttf"] = _obfuscate_font(
        (FONT_ROOT / "NotoSansKR-Medium.otf").read_bytes(), regular_key
    )
    entries["word/fonts/NotoSansKR-Bold.odttf"] = _obfuscate_font(
        (FONT_ROOT / "NotoSansKR-Bold.otf").read_bytes(), bold_key
    )

    with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w") as destination:
            for name in sorted(entries):
                member = zipfile.ZipInfo(name, date_time=DOCX_ZIP_TIMESTAMP)
                member.compress_type = zipfile.ZIP_DEFLATED
                member.create_system = 0
                member.external_attr = 0
                member.internal_attr = 0
                destination.writestr(
                    member,
                    entries[name],
                    compress_type=zipfile.ZIP_DEFLATED,
                    compresslevel=9,
                )
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def build_guide(source: Path, output: Path) -> None:
    _validate_fonts()
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    core = document.core_properties
    core.title = TITLE
    core.subject = SUBJECT
    core.version = VERSION
    core.category = STYLE_CATEGORY
    core.author = "IROA.AI"
    core.last_modified_by = "IROA.AI"
    core.keywords = "IROA.AI, Brand Identity, Track A, v1.0, 2026-08-27"
    core.comments = "Source: docs/brand/IROA_BI_GUIDE_KO.md; official BI: Track A"
    timestamp = datetime(2026, 8, 27, tzinfo=timezone.utc)
    core.created = timestamp
    core.modified = timestamp

    _configure_section(document)
    _configure_styles(document)
    _configure_headers_and_footers(document)
    bullet_num_id, decimal_num_id = _list_numbering(document)
    body_start = _add_cover(document, lines, bullet_num_id)
    _add_markdown_body(document, lines, body_start, bullet_num_id, decimal_num_id)
    _enable_field_updates(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        document.save(temporary)
        _embed_pinned_fonts(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the IROA.AI BI v1.0 DOCX distribution guide")
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build_guide(args.source.resolve(), args.output.resolve())
    print(f"built {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
