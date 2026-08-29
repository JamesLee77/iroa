#!/usr/bin/env python3
"""Build the English IROA whitepaper DOCX from its reviewed Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/whitepaper/IROA_WHITEPAPER_EN.md"
OUTPUT = ROOT / "docs/whitepaper/exports/IROA_WHITEPAPER_EN.docx"
BRAND_BLUE = "3268A8"
INK = RGBColor(31, 43, 55)
MUTED = RGBColor(82, 96, 110)


def clean_inline(value: str) -> str:
    value = re.sub(r"!\[([^]]*)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value)
    value = value.replace("**", "").replace("__", "").replace("`", "")
    return value.strip()


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = properties.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        properties.append(node)
    node.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_cell_text(cell, value: str, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(clean_inline(value))
    run.bold = header
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(255, 255, 255) if header else INK
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        shade(cell, BRAND_BLUE)


def add_table(document: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            value = values[column_index] if column_index < len(values) else ""
            set_cell_text(table.cell(row_index, column_index), value, row_index == 0)
    document.add_paragraph()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_image(document: Document, markdown_line: str) -> None:
    match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", markdown_line.strip())
    if not match:
        return
    alt, raw_path = match.groups()
    image_path = (SOURCE.parent / raw_path).resolve()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(clean_inline(alt))


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    for style_name in ["Normal", "Body Text"]:
        style = document.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10)
        style.font.color.rgb = INK
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
    for level, size in [(1, 26), (2, 19), (3, 13)]:
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BRAND_BLUE if level < 3 else "234A75")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(7)
    document.styles["Caption"].font.name = "Aptos"
    document.styles["Caption"].font.size = Pt(8.5)
    document.styles["Caption"].font.color.rgb = MUTED
    header = section.header.paragraphs[0]
    header.text = "IROA.AI  /  WHITEPAPER  /  ENGLISH EDITION"
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = MUTED
    add_page_number(section.footer.paragraphs[0])


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)
    document.core_properties.title = "IROA.AI Whitepaper — English Edition"
    document.core_properties.subject = "Inclusive Real-world Orchestration Agent"
    document.core_properties.author = "IROA.AI"
    document.core_properties.keywords = "IROA, accessibility, secure execution, Node, token economy"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(108)
    run = title.add_run("IROA.AI")
    run.font.name = "Aptos Display"
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BRAND_BLUE)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("WHITEPAPER · ENGLISH EDITION").bold = True
    edition = document.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    edition.add_run("Final edition · v1.0 · 2026-08-21\nRepublic of Korea first\nKorean edition controls")
    document.add_page_break()

    contents = document.add_paragraph("Contents", style="Heading 1")
    contents.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for line in lines:
        match = re.match(r"^## (\d+\. .+)$", line)
        if match:
            document.add_paragraph(match.group(1), style="List Number")
    document.add_page_break()

    index = 1
    in_code = False
    code_lines: list[str] = []
    first_chapter = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("~~~"):
            if in_code:
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["No Spacing"]
                paragraph.paragraph_format.left_indent = Inches(0.25)
                run = paragraph.add_run("\n".join(code_lines))
                run.font.name = "Aptos Mono"
                run.font.size = Pt(8.5)
                shade_cell = OxmlElement("w:shd")
                shade_cell.set(qn("w:fill"), "EEF3F8")
                paragraph._p.get_or_add_pPr().append(shade_cell)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped or stripped == "---":
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if stripped.startswith("!["):
            add_image(document, stripped)
            index += 1
            continue
        chapter = re.match(r"^## (\d+\. .+)$", stripped)
        if chapter:
            if not first_chapter:
                document.add_page_break()
            first_chapter = False
            document.add_paragraph(clean_inline(chapter.group(1)), style="Heading 1")
            index += 1
            continue
        if stripped.startswith("# "):
            index += 1
            continue
        if stripped.startswith("## "):
            document.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
        elif stripped.startswith("### "):
            document.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
        elif re.match(r"^\d+\. ", stripped):
            document.add_paragraph(clean_inline(re.sub(r"^\d+\. ", "", stripped)), style="List Number")
        elif stripped.startswith("- "):
            document.add_paragraph(clean_inline(stripped[2:]), style="List Bullet")
        elif stripped.startswith("> "):
            paragraph = document.add_paragraph(clean_inline(stripped[2:]))
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.color.rgb = MUTED
        else:
            document.add_paragraph(clean_inline(stripped))
        index += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
