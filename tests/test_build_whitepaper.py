import hashlib
from itertools import groupby
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
import pymupdf
import pytest

from scripts.build_whitepaper import build


ROOT = Path(__file__).parents[1]
OFFICIAL_SOURCE = ROOT / "docs/whitepaper/MODUA_WHITEPAPER_KO.md"
LIST_GROUP_SIZES = [
    5,
    5,
    3,
    3,
    3,
    3,
    3,
    3,
    4,
    4,
    4,
    4,
    5,
    3,
    4,
    3,
    5,
    4,
    4,
    4,
    4,
    4,
    4,
    4,
    4,
    4,
    6,
    2,
    3,
    2,
    1,
    2,
]


@pytest.fixture(scope="module")
def official_artifacts(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path, Path]:
    output_dir = tmp_path_factory.mktemp("official-whitepaper")
    docx_path, pdf_path = build(OFFICIAL_SOURCE, output_dir)
    return OFFICIAL_SOURCE, docx_path, pdf_path


def _compact(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _assert_in_order(expected: list[str], actual: str) -> None:
    cursor = 0
    for value in expected:
        cursor = actual.find(value, cursor)
        assert cursor >= 0, f"missing or out of order: {value}"
        cursor += len(value)


def _source_headings(source: str) -> list[tuple[int, str]]:
    headings: list[tuple[int, str]] = []
    for line in source.splitlines():
        match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if match:
            headings.append((len(match.group(1)), match.group(2).strip()))
    return headings


def _source_table_cells(source: str) -> list[str]:
    cells: list[str] = []
    for line in source.splitlines():
        if not line.startswith("|"):
            continue
        values = [value.strip() for value in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            continue
        cells.extend(re.sub(r"[*_`]", "", value) for value in values)
    return cells


def _human_visible_markdown(text: str) -> str:
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(`+)(.*?)\1", r"\2", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    return re.sub(r"(?<!\*)\*([^*]+)\*", r"\1", text).strip()


def _source_body_items_and_list_groups(
    source: str,
) -> tuple[list[str], list[tuple[str, list[str]]]]:
    body_items: list[str] = []
    list_groups: list[tuple[str, list[str]]] = []
    prose_lines: list[str] = []
    active_list_index: int | None = None

    def flush_prose() -> None:
        if prose_lines:
            body_items.append(_human_visible_markdown(" ".join(prose_lines)))
            prose_lines.clear()

    for line in [*source.splitlines(), ""]:
        if not line.strip():
            flush_prose()
            active_list_index = None
            continue
        if re.match(r"^#{1,6}\s+", line) or line.startswith("|"):
            flush_prose()
            active_list_index = None
            continue

        list_match = re.match(r"^(?:([-+*])|(\d+)\.)\s+(.+)$", line)
        if list_match:
            flush_prose()
            list_kind = "bullet" if list_match.group(1) else "decimal"
            if active_list_index is None or list_groups[active_list_index][0] != list_kind:
                list_groups.append((list_kind, []))
                active_list_index = len(list_groups) - 1
            item = _human_visible_markdown(list_match.group(3))
            list_groups[active_list_index][1].append(item)
            body_items.append(item)
            continue

        active_list_index = None
        prose_lines.append(re.sub(r"^>\s?", "", line))

    return body_items, list_groups


def _docx_text(document: DocumentObject) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + cells)


def _pdf_body_text(pdf: pymupdf.Document) -> str:
    words: list[str] = []
    for page in pdf:
        words.extend(
            word[4]
            for word in page.get_text("words")
            if word[1] >= 40
            and word[3] <= 800
            and not re.fullmatch(r"(?:•|\d+\.)", word[4])
        )
    return _compact(" ".join(words))


def test_build_creates_matching_docx_and_pdf(tmp_path: Path) -> None:
    source = tmp_path / "sample.md"
    source.write_text(
        "# MODUA\n\n## 1. 핵심 선언\n\n모두의 일상을 돕는 AI\n\n| 항목 | 결정 |\n|---|---|\n| 브랜드 | MODUA |",
        encoding="utf-8",
    )
    docx_path, pdf_path = build(source, tmp_path / "exports")
    assert docx_path.stat().st_size > 0
    assert pdf_path.stat().st_size > 0
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    with pymupdf.open(pdf_path) as pdf:
        pdf_text = "\n".join(page.get_text() for page in pdf)
    assert "MODUA" in docx_text and "1. 핵심 선언" in docx_text
    assert "MODUA" in pdf_text and "1. 핵심 선언" in pdf_text


def test_official_source_content_survives_in_both_outputs(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    source_path, docx_path, pdf_path = official_artifacts
    source = source_path.read_text(encoding="utf-8")
    headings = [text for _, text in _source_headings(source)]
    table_cells = _source_table_cells(source)
    body_items, list_groups = _source_body_items_and_list_groups(source)
    assert len(body_items) == 210
    assert sum(len(items) for _, items in list_groups) == 116

    document = Document(docx_path)
    docx_text = _compact(_docx_text(document))
    docx_headings = [
        _compact(paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    docx_body_items = [
        _compact(paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.style.name not in {"Heading 1", "Heading 2", "Heading 3"}
    ]
    with pymupdf.open(pdf_path) as pdf:
        pdf_text = _compact("\n".join(page.get_text() for page in pdf))
        pdf_body_text = _pdf_body_text(pdf)

    assert docx_headings == [_compact(text) for text in headings]
    assert docx_body_items == [_compact(text) for text in body_items]
    _assert_in_order([_compact(text) for text in headings], pdf_text)
    _assert_in_order([_compact(text) for text in body_items], pdf_body_text)
    for value in table_cells:
        expected = _compact(value)
        assert expected in docx_text
        assert expected in pdf_text


def test_official_docx_uses_accessible_a4_publication_structure(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    source_path, docx_path, _ = official_artifacts
    document = Document(docx_path)
    section = document.sections[0]
    expected_headings = _source_headings(source_path.read_text(encoding="utf-8"))
    actual_headings = [
        (int(paragraph.style.name[-1]), paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]

    assert actual_headings == expected_headings
    for actual, expected in [
        (section.page_width.mm, 210),
        (section.page_height.mm, 297),
        (section.top_margin.mm, 20),
        (section.left_margin.mm, 18),
        (section.right_margin.mm, 18),
        (section.bottom_margin.mm, 22),
        (section.header_distance.mm, 12.5),
        (section.footer_distance.mm, 12.5),
    ]:
        assert actual == pytest.approx(expected, abs=0.02)

    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name in {"MODUA Bullet", "MODUA Number"}
    ]
    num_ids = [
        paragraph._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val"))
        for paragraph in list_paragraphs
    ]
    groups = [(num_id, len(list(items))) for num_id, items in groupby(num_ids)]
    assert [size for _, size in groups] == LIST_GROUP_SIZES
    assert len({num_id for num_id, _ in groups}) == len(groups)

    assert section.different_first_page_header_footer
    assert section.first_page_header.paragraphs[0].text == ""
    assert section.first_page_footer.paragraphs[0].text == ""
    assert section.header.paragraphs[0].text == "MODUA WHITEPAPER · KOREAN REVIEW EDITION"
    page_fields = section.footer._element.xpath(".//w:instrText")
    assert [field.text.strip() for field in page_fields] == ["PAGE"]

    assert len(document.tables) == 5
    for table in document.tables:
        grid_widths = [
            int(column.get(qn("w:w")))
            for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
        ]
        assert sum(grid_widths) == 9745
        table_properties = table._tbl.tblPr
        assert table_properties.find(qn("w:tblW")).get(qn("w:w")) == "9745"
        assert table_properties.find(qn("w:tblInd")).get(qn("w:w")) == "120"
        assert table_properties.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
        assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        assert all(not row._tr.xpath("./w:trPr/w:trHeight") for row in table.rows)
        for row in table.rows:
            for cell, expected_width in zip(row.cells, grid_widths, strict=True):
                properties = cell._tc.get_or_add_tcPr()
                assert int(properties.find(qn("w:tcW")).get(qn("w:w"))) == expected_width
                assert properties.find(qn("w:vAlign")).get(qn("w:val")) == "center"
                margins = properties.find(qn("w:tcMar"))
                for side, expected_margin in {
                    "top": "80",
                    "bottom": "80",
                    "start": "120",
                    "end": "120",
                }.items():
                    assert margins.find(qn(f"w:{side}")).get(qn("w:w")) == expected_margin


def test_official_docx_styles_and_numbering_match_accessible_contract(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    source_path, docx_path, _ = official_artifacts
    document = Document(docx_path)

    normal = document.styles["Normal"]
    assert normal.font.size.pt == pytest.approx(12)
    assert normal.element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Sans KR"
    assert normal.paragraph_format.line_spacing == pytest.approx(1.333, abs=0.001)
    assert normal.paragraph_format.widow_control is True
    normal_spacing = normal.element.pPr.find(qn("w:spacing"))
    assert normal_spacing.get(qn("w:line")) == "320"
    assert normal_spacing.get(qn("w:lineRule")) == "auto"
    assert normal.element.pPr.find(qn("w:widowControl")) is not None

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10, "0"),
        "Heading 2": (13, "2E74B5", 12, 6, "1"),
        "Heading 3": (12, "1F4D78", 8, 4, "2"),
    }
    for name, (size, color, before, after, outline_level) in heading_tokens.items():
        style = document.styles[name]
        assert style.font.size.pt == pytest.approx(size)
        assert str(style.font.color.rgb) == color
        assert style.element.rPr.rFonts.get(qn("w:eastAsia")) == "Noto Sans KR"
        assert style.paragraph_format.space_before.pt == pytest.approx(before)
        assert style.paragraph_format.space_after.pt == pytest.approx(after)
        assert style.paragraph_format.keep_with_next is True
        assert style.paragraph_format.keep_together is True
        assert style.element.pPr.find(qn("w:keepNext")) is not None
        assert style.element.pPr.find(qn("w:keepLines")) is not None
        assert style.element.pPr.find(qn("w:outlineLvl")).get(qn("w:val")) == outline_level

    _, source_list_groups = _source_body_items_and_list_groups(
        source_path.read_text(encoding="utf-8")
    )
    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name in {"MODUA Bullet", "MODUA Number"}
    ]
    num_ids = [
        paragraph._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val"))
        for paragraph in list_paragraphs
    ]
    generated_groups = [
        (num_id, list(items))
        for num_id, items in groupby(
            zip(num_ids, list_paragraphs), key=lambda pair: pair[0]
        )
    ]
    numbering = document.part.numbering_part.element
    generated_formats: list[str] = []
    for num_id, items in generated_groups:
        num = next(
            node
            for node in numbering.findall(qn("w:num"))
            if node.get(qn("w:numId")) == num_id
        )
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        abstract = next(
            node
            for node in numbering.findall(qn("w:abstractNum"))
            if node.get(qn("w:abstractNumId")) == abstract_id
        )
        generated_formats.append(
            abstract.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val"))
        )
        expected_style = "MODUA Bullet" if generated_formats[-1] == "bullet" else "MODUA Number"
        assert {paragraph.style.name for _, paragraph in items} == {expected_style}

    assert generated_formats == [kind for kind, _ in source_list_groups]
    assert [len(items) for _, items in generated_groups] == LIST_GROUP_SIZES


def test_official_docx_meets_repeatable_high_roi_accessibility_gate(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    _, docx_path, _ = official_artifacts
    document = Document(docx_path)
    heading_levels = [
        int(paragraph.style.name[-1])
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
    ]

    assert heading_levels[0] == 1
    assert all(current <= previous + 1 for previous, current in zip(heading_levels, heading_levels[1:]))
    assert all(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader") for table in document.tables)
    assert not document.inline_shapes
    assert not document._element.xpath(".//w:drawing")
    assert not document._element.xpath(".//w:pict")
    assert not document._element.xpath(".//w:hyperlink")


def test_official_pdf_has_complete_table_header_tag_graph(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    _, _, pdf_path = official_artifacts
    with pymupdf.open(pdf_path) as pdf:
        catalog = pdf.pdf_catalog()
        mark_kind, mark_info = pdf.xref_get_key(catalog, "MarkInfo")
        tree_kind, tree_reference = pdf.xref_get_key(catalog, "StructTreeRoot")
        assert mark_kind == "dict" and "/Marked true" in mark_info
        assert tree_kind == "xref" and tree_reference.endswith(" R")

        table_header_ids: list[str] = []
        header_references: list[str] = []
        for xref in range(1, pdf.xref_length()):
            if pdf.xref_get_key(xref, "S") == ("name", "/TH"):
                id_kind, header_id = pdf.xref_get_key(xref, "ID")
                assert id_kind == "string" and header_id.startswith("modua-")
                table_header_ids.append(header_id)
            reference_kind, references = pdf.xref_get_key(xref, "A/Headers")
            if reference_kind == "array":
                values = re.findall(r"\(([^()]*)\)", references)
                assert values
                header_references.extend(values)

        assert table_header_ids
        assert len(table_header_ids) == len(set(table_header_ids))
        assert header_references
        assert set(header_references) == set(table_header_ids)
        assert len(pdf) == 30
        for page in pdf:
            assert page.rect.width == pytest.approx(595.276, abs=0.02)
            assert page.rect.height == pytest.approx(841.89, abs=0.02)
            for x0, y0, x1, y1, *_ in page.get_text("words"):
                assert 0 <= x0 <= x1 <= page.rect.width
                assert 0 <= y0 <= y1 <= page.rect.height


def test_official_pdf_has_running_header_and_ordered_page_footers(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    _, _, pdf_path = official_artifacts
    expected_header = "MODUA WHITEPAPER · KOREAN REVIEW EDITION"
    footer_numbers: list[str] = []

    with pymupdf.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            words = page.get_text("words")
            header_words = sorted((word for word in words if word[1] < 40), key=lambda word: word[0])
            footer_words = sorted((word for word in words if word[3] > 800), key=lambda word: word[0])
            if page_number == 1:
                assert header_words == []
                assert footer_words == []
                continue

            assert " ".join(word[4] for word in header_words) == expected_header
            assert [word[4] for word in footer_words] == [str(page_number)]
            assert all(800 < word[1] < word[3] <= page.rect.height for word in footer_words)
            assert all(abs((word[0] + word[2]) / 2 - page.rect.width / 2) < 1 for word in footer_words)
            footer_numbers.extend(word[4] for word in footer_words)

    assert footer_numbers == [str(page_number) for page_number in range(2, 31)]


def test_separate_ordered_lists_restart_numbering(tmp_path: Path) -> None:
    source = tmp_path / "lists.md"
    source.write_text(
        "# 목록\n\n1. 첫째\n2. 둘째\n\n## 다음 목록\n\n1. 다시 첫째\n2. 다시 둘째",
        encoding="utf-8",
    )
    docx_path, _ = build(source, tmp_path / "exports")
    document = Document(docx_path)
    numbered = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "MODUA Number"]
    num_ids = [
        paragraph._p.xpath("./w:pPr/w:numPr/w:numId")[0].get(qn("w:val"))
        for paragraph in numbered
    ]
    assert num_ids[:2] == [num_ids[0], num_ids[0]]
    assert num_ids[2:] == [num_ids[2], num_ids[2]]
    assert num_ids[0] != num_ids[2]


def test_two_column_table_reserves_width_for_korean_row_labels(tmp_path: Path) -> None:
    source = tmp_path / "table.md"
    source.write_text(
        "# 표\n\n| 위험 | 대응과 중단 경계 |\n|---|---|\n"
        "| 도움 제공자 담합 | " + "긴 대응 설명 " * 10 + "|",
        encoding="utf-8",
    )
    docx_path, _ = build(source, tmp_path / "exports")
    table = Document(docx_path).tables[0]
    widths = [
        int(column.get(qn("w:w")))
        for column in table._tbl.tblGrid.findall(qn("w:gridCol"))
    ]
    assert widths[0] >= 2924


def test_unsupported_top_level_markdown_fails_before_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "unsupported.md"
    source.write_text("# 지원 범위\n\n---\n\n누락되면 안 되는 내용", encoding="utf-8")
    output_dir = tmp_path / "exports"

    with pytest.raises(ValueError, match=r"unsupported top-level Markdown/HTML node: hr"):
        build(source, output_dir)

    assert not list(output_dir.glob("MODUA_WHITEPAPER_KO.*"))


def test_build_is_byte_reproducible(tmp_path: Path) -> None:
    first_docx, first_pdf = build(OFFICIAL_SOURCE, tmp_path / "first")
    second_docx, second_pdf = build(OFFICIAL_SOURCE, tmp_path / "second")
    digest = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
    assert digest(first_docx) == digest(second_docx)
    assert digest(first_pdf) == digest(second_pdf)


def test_committed_and_generated_docx_declare_korean_language_metadata(
    official_artifacts: tuple[Path, Path, Path],
) -> None:
    _, generated_docx, _ = official_artifacts
    committed_docx = ROOT / "docs/whitepaper/exports/MODUA_WHITEPAPER_KO.docx"
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    word = f"{{{word_namespace}}}"

    for docx_path in (committed_docx, generated_docx):
        with zipfile.ZipFile(docx_path) as archive:
            settings = ET.fromstring(archive.read("word/settings.xml"))
            theme_language = settings.find(f".//{word}themeFontLang")
            assert theme_language is not None
            assert theme_language.get(f"{word}val") == "ko-KR"
            assert theme_language.get(f"{word}eastAsia") == "ko-KR"

            for part_name in ("word/styles.xml", "word/stylesWithEffects.xml"):
                styles = ET.fromstring(archive.read(part_name))
                default_language = styles.find(
                    f"./{word}docDefaults/{word}rPrDefault/{word}rPr/{word}lang"
                )
                assert default_language is not None
                assert default_language.get(f"{word}val") == "ko-KR"
                assert default_language.get(f"{word}eastAsia") == "ko-KR"

            text_parts = [
                name
                for name in archive.namelist()
                if name == "word/document.xml"
                or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            ]
            assert text_parts
            for part_name in text_parts:
                root = ET.fromstring(archive.read(part_name))
                text_runs = [
                    run
                    for run in root.iter(f"{word}r")
                    if run.find(f"{word}t") is not None
                    or run.find(f"{word}instrText") is not None
                ]
                for run in text_runs:
                    language = run.find(f"{word}rPr/{word}lang")
                    assert language is not None, f"missing run language in {part_name}"
                    assert language.get(f"{word}val") == "ko-KR"
                    assert language.get(f"{word}eastAsia") == "ko-KR"

            for part_name in (
                name
                for name in archive.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            ):
                root = ET.fromstring(archive.read(part_name))
                for element_name in ("lang", "themeFontLang"):
                    for language in root.iter(f"{word}{element_name}"):
                        values = {
                            language.get(f"{word}val"),
                            language.get(f"{word}eastAsia"),
                        }
                        assert not values.intersection({"en-US", "ja-JP"})
