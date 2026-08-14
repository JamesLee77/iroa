from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import hashlib
import os
import re
import zipfile

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from markdown import markdown
import pymupdf
from weasyprint import HTML


TITLE = "MODUA Whitepaper"
SUBJECT = "모두의 일상을 돕는 AI 생활자립·상호지원 생태계"
OUTPUT_STEM = "MODUA_WHITEPAPER_KO"
FIXED_TIMESTAMP = datetime(2026, 8, 14, tzinfo=timezone.utc)
DOCUMENT_LANGUAGE = "ko-KR"

# narrative_proposal, with the named modua_a4_publication and
# large_text_a11y overrides applied explicitly.
BODY_FONT = "Calibri"
KOREAN_FONT = "Noto Sans KR"
BODY_SIZE_PT = 12
PAGE_USABLE_DXA = 9865  # A4 width minus 18 mm left/right margins.
TABLE_INDENT_DXA = 120
TABLE_WIDTH_DXA = PAGE_USABLE_DXA - TABLE_INDENT_DXA
TABLE_CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
HEADER_FILL = "F4F6F9"
ACCENT = RGBColor(46, 116, 181)
ACCENT_DARK = RGBColor(31, 77, 120)
TEXT = RGBColor(23, 23, 23)
MUTED = RGBColor(92, 104, 116)


CSS = """
@page {
  size: A4 portrait;
  margin: 20mm 18mm 22mm;
  @top-center {
    content: "MODUA WHITEPAPER · KOREAN REVIEW EDITION";
    font-family: 'Noto Sans CJK KR', 'Apple SD Gothic Neo', sans-serif;
    font-size: 8.5pt;
    color: #66717d;
  }
  @bottom-center {
    content: counter(page);
    font-family: 'Noto Sans CJK KR', 'Apple SD Gothic Neo', sans-serif;
    font-size: 9pt;
    color: #66717d;
  }
}
@page:first {
  @top-center { content: none; }
  @bottom-center { content: none; }
}
html { font-size: 16pt; }
body {
  font-family: 'Noto Sans CJK KR', 'Apple SD Gothic Neo', 'Malgun Gothic', sans-serif;
  color: #171717;
  font-size: 16pt;
  line-height: 1.55;
  margin: 0;
  overflow-wrap: anywhere;
}
.cover {
  min-height: 235mm;
  box-sizing: border-box;
  padding-top: 54mm;
  text-align: center;
  page-break-after: always;
}
.cover h1 {
  color: #203748;
  font-size: 30pt;
  line-height: 1.18;
  margin: 0 0 10pt;
  page-break-before: avoid;
}
.cover h2 {
  color: #2b5163;
  font-size: 15pt;
  line-height: 1.45;
  margin: 0 0 32pt;
}
.cover blockquote {
  border: 0;
  color: #9a7420;
  font-size: 11pt;
  font-weight: 600;
  margin: 0;
  padding: 0;
}
h1, h2, h3 { break-after: avoid; page-break-after: avoid; }
h1 { color: #2e74b5; font-size: 28pt; line-height: 1.2; margin: 18pt 0 10pt; }
h2 { color: #2e74b5; font-size: 22pt; line-height: 1.25; margin: 20pt 0 8pt; }
h3 { color: #1f4d78; font-size: 18pt; line-height: 1.3; margin: 14pt 0 6pt; }
p { margin: 0 0 10pt; orphans: 2; widows: 2; }
ul, ol { margin: 0 0 10pt; padding-left: 1.65em; }
li { margin: 0 0 4pt; break-inside: avoid; }
blockquote {
  border-left: 4px solid #334e68;
  color: #334e68;
  margin: 12pt 0;
  padding: 3pt 0 3pt 12pt;
  break-inside: avoid;
}
table {
  width: 100%;
  border-collapse: collapse;
  table-layout: fixed;
  margin: 12pt 0 14pt;
}
thead { display: table-header-group; }
tr { break-inside: avoid; page-break-inside: avoid; }
th, td {
  border: 0.75pt solid #aab3bd;
  padding: 7pt;
  text-align: left;
  vertical-align: middle;
  line-height: 1.35;
  overflow-wrap: anywhere;
}
th { background: #f4f6f9; font-weight: 700; }
code {
  font-family: 'D2Coding', 'Apple SD Gothic Neo', monospace;
  font-size: 0.88em;
  white-space: normal;
}
a { color: #1f4d78; text-decoration: underline; }
"""


def _set_style_font(style, *, size: float, color: RGBColor | None = None) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    fonts.set(qn("w:eastAsia"), KOREAN_FONT)
    fonts.set(qn("w:cs"), BODY_FONT)


def _set_language(element) -> None:
    language = element.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        element.append(language)
    language.set(qn("w:val"), DOCUMENT_LANGUAGE)
    language.set(qn("w:eastAsia"), DOCUMENT_LANGUAGE)
    language.attrib.pop(qn("w:bidi"), None)


def _configure_language_metadata(document: DocumentObject) -> None:
    settings = document.settings.element
    theme_language = settings.find(qn("w:themeFontLang"))
    if theme_language is None:
        theme_language = OxmlElement("w:themeFontLang")
        settings.append(theme_language)
    theme_language.set(qn("w:val"), DOCUMENT_LANGUAGE)
    theme_language.set(qn("w:eastAsia"), DOCUMENT_LANGUAGE)
    theme_language.attrib.pop(qn("w:bidi"), None)

    default_run_properties = document.styles.element.xpath(
        "./w:docDefaults/w:rPrDefault/w:rPr"
    )
    if not default_run_properties:
        raise RuntimeError("DOCX template is missing default run properties")
    _set_language(default_run_properties[0])

    roots = [document.element]
    for section in document.sections:
        roots.extend(
            [
                section.header._element,
                section.first_page_header._element,
                section.footer._element,
                section.first_page_footer._element,
            ]
        )
    seen_roots: set[int] = set()
    for root in roots:
        if id(root) in seen_roots:
            continue
        seen_roots.add(id(root))
        for run in root.xpath(".//w:r"):
            if run.find(qn("w:t")) is None and run.find(qn("w:instrText")) is None:
                continue
            _set_language(run.get_or_add_rPr())


def _set_paragraph_border(paragraph, *, side: str, color: str, size: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "5")
    border.set(qn("w:color"), color)
    borders.append(border)


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles

    normal = styles["Normal"]
    _set_style_font(normal, size=BODY_SIZE_PT, color=TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, ACCENT, 18, 10),
        "Heading 2": (13, ACCENT, 12, 6),
        "Heading 3": (12, ACCENT_DARK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        _set_style_font(style, size=size, color=color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    quote = styles["Quote"]
    _set_style_font(quote, size=BODY_SIZE_PT, color=ACCENT_DARK)
    quote.font.italic = False
    quote.paragraph_format.left_indent = Mm(5)
    quote.paragraph_format.right_indent = Mm(2)
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(10)
    quote.paragraph_format.line_spacing = 1.333

    for name in ("MODUA Bullet", "MODUA Number"):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]
        style.base_style = normal
        _set_style_font(style, size=BODY_SIZE_PT, color=TEXT)
        style.paragraph_format.left_indent = Mm(9.525)
        style.paragraph_format.first_line_indent = Mm(-4.928)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208
        style.paragraph_format.widow_control = True

    if "MODUA Table Text" not in styles:
        styles.add_style("MODUA Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text = styles["MODUA Table Text"]
    table_text.base_style = normal
    _set_style_font(table_text, size=BODY_SIZE_PT, color=TEXT)
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.15


def _add_numbering_definition(document: DocumentObject, *, ordered: bool) -> int:
    numbering = document.part.numbering_part.element
    existing_abstract = [
        int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(existing_abstract, default=-1) + 1
    existing_num = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_num, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "279")
    ppr.append(ind)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)


def _append_inline(paragraph, node, *, bold: bool = False, italic: bool = False, code: bool = False) -> None:
    if isinstance(node, NavigableString):
        if not str(node):
            return
        run = paragraph.add_run(str(node))
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = KOREAN_FONT
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), KOREAN_FONT)
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), KOREAN_FONT)
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
            run.font.size = Pt(10.5)
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    next_bold = bold or node.name in {"b", "strong"}
    next_italic = italic or node.name in {"i", "em"}
    next_code = code or node.name == "code"
    for child in node.children:
        _append_inline(paragraph, child, bold=next_bold, italic=next_italic, code=next_code)


def _add_html_paragraph(document: DocumentObject, node: Tag, *, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    for child in node.children:
        _append_inline(paragraph, child)
    return paragraph


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in TABLE_CELL_MARGINS_DXA.items():
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    table_width = tbl_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "AAB3BD")
        borders.append(border)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            _set_cell_width(cell, width)
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _column_widths(rows: list[Tag], column_count: int) -> list[int]:
    weights = [1] * column_count
    for row in rows:
        for index, cell in enumerate(row.find_all(["th", "td"], recursive=False)):
            weights[index] = max(weights[index], min(len(cell.get_text(" ", strip=True)), 60))
    floor = max(1, int(TABLE_WIDTH_DXA * 0.16))
    remaining = TABLE_WIDTH_DXA - floor * column_count
    if remaining < 0:
        base = TABLE_WIDTH_DXA // column_count
        widths = [base] * column_count
    else:
        total_weight = sum(weights)
        widths = [floor + int(remaining * weight / total_weight) for weight in weights]
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    if column_count == 2:
        minimum_label_width = (TABLE_WIDTH_DXA * 30 + 99) // 100
        if widths[0] < minimum_label_width:
            widths[1] -= minimum_label_width - widths[0]
            widths[0] = minimum_label_width
    return widths


def _add_table(document: DocumentObject, node: Tag) -> None:
    rows = node.find_all("tr")
    if not rows:
        return
    column_count = max(len(row.find_all(["th", "td"], recursive=False)) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    for row_index, source_row in enumerate(rows):
        target_row = table.add_row()
        source_cells = source_row.find_all(["th", "td"], recursive=False)
        for column_index, source_cell in enumerate(source_cells):
            target = target_row.cells[column_index]
            paragraph = target.paragraphs[0]
            paragraph.style = document.styles["MODUA Table Text"]
            for child in source_cell.children:
                _append_inline(paragraph, child, bold=row_index == 0)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), HEADER_FILL)
                target._tc.get_or_add_tcPr().append(shading)
        if row_index == 0:
            _set_repeat_table_header(target_row)
    _set_table_geometry(table, _column_widths(rows, column_count))
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _append_docx(document: DocumentObject, soup: BeautifulSoup) -> None:
    in_cover = True
    cover_title_seen = False
    for node in soup.children:
        if not isinstance(node, Tag):
            continue
        text = node.get_text(" ", strip=True)
        if in_cover and node.name == "h2" and re.match(r"^1\.\s", text):
            document.add_page_break()
            in_cover = False

        if node.name in {"h1", "h2", "h3"}:
            level = int(node.name[1])
            paragraph = _add_html_paragraph(document, node, style=f"Heading {level}")
            if in_cover:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                if node.name == "h1":
                    cover_title_seen = True
                    paragraph.paragraph_format.space_before = Pt(108)
                    paragraph.paragraph_format.space_after = Pt(10)
                    for run in paragraph.runs:
                        run.font.size = Pt(30)
                        run.font.color.rgb = RGBColor(32, 55, 72)
                elif node.name == "h2" and cover_title_seen:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(32)
                    for run in paragraph.runs:
                        run.font.size = Pt(15)
                        run.font.color.rgb = RGBColor(43, 81, 99)
        elif node.name == "p":
            _add_html_paragraph(document, node)
        elif node.name in {"ul", "ol"}:
            style = "MODUA Bullet" if node.name == "ul" else "MODUA Number"
            num_id = _add_numbering_definition(document, ordered=node.name == "ol")
            for item in node.find_all("li", recursive=False):
                paragraph = _add_html_paragraph(document, item, style=style)
                _apply_numbering(paragraph, num_id)
        elif node.name == "blockquote":
            paragraph = _add_html_paragraph(document, node, style="Quote")
            if in_cover:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.left_indent = Mm(0)
                paragraph.paragraph_format.right_indent = Mm(0)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(154, 116, 32)
            else:
                _set_paragraph_border(paragraph, side="left", color="334E68", size="24")
        elif node.name == "table":
            _add_table(document, node)


def _append_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "2"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def _configure_page(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = 0
    section.top_margin = Mm(20)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.bottom_margin = Mm(22)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header.paragraphs[0]
    first_header.text = ""
    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.text = ""

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("MODUA WHITEPAPER · KOREAN REVIEW EDITION")
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.space_before = Pt(0)
    _append_page_field(footer)


def _normalize_docx_archive(path: Path) -> None:
    normalized = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for info in source.infolist():
            replacement = zipfile.ZipInfo(info.filename, (1980, 1, 1, 0, 0, 0))
            replacement.compress_type = zipfile.ZIP_DEFLATED
            replacement.external_attr = info.external_attr
            replacement.create_system = info.create_system
            data = source.read(info.filename)
            if info.filename in {"word/styles.xml", "word/stylesWithEffects.xml"}:
                data, replacements = re.subn(
                    rb"<w:lang\b[^>]*/>",
                    b'<w:lang w:val="ko-KR" w:eastAsia="ko-KR"/>',
                    data,
                )
                if replacements == 0:
                    raise RuntimeError(f"missing default language metadata: {info.filename}")
            target.writestr(replacement, data)
    os.replace(normalized, path)


def _normalize_pdf(path: Path, identifier: str) -> None:
    """Remove runtime-derived tagged-table IDs from a WeasyPrint PDF."""
    normalized = path.with_suffix(".normalized.pdf")
    with pymupdf.open(path) as pdf:
        stable_ids: dict[str, str] = {}
        for xref in range(1, pdf.xref_length()):
            kind, value = pdf.xref_get_key(xref, "ID")
            if kind == "string":
                stable_ids[value] = f"modua-{xref}"

        for xref in range(1, pdf.xref_length()):
            kind, value = pdf.xref_get_key(xref, "ID")
            if kind == "string":
                pdf.xref_set_key(xref, "ID", f"({stable_ids[value]})")

            kind, value = pdf.xref_get_key(xref, "A/Headers")
            if kind == "array":
                for runtime_id, stable_id in stable_ids.items():
                    value = value.replace(f"({runtime_id})", f"({stable_id})")
                pdf.xref_set_key(xref, "A/Headers", value)

        pdf.save(
            normalized,
            garbage=4,
            clean=True,
            deflate=True,
            deflate_images=True,
            deflate_fonts=True,
            no_new_id=True,
        )

    data = normalized.read_bytes()
    deterministic_trailer = (
        b"/ID[(" + identifier.encode("ascii") + b")(" + identifier[:32].encode("ascii") + b")]"
    )
    data, replacements = re.subn(
        rb"/ID\s*\[\s*\([^)]*\)\s*\([^)]*\)\s*\]",
        deterministic_trailer,
        data,
    )
    if replacements != 1:
        raise RuntimeError(f"expected one PDF trailer ID, found {replacements}")
    normalized.write_bytes(data)
    os.replace(normalized, path)


def _cover_html(soup: BeautifulSoup) -> str:
    cover: list[str] = []
    body: list[str] = []
    in_cover = True
    for node in soup.children:
        if isinstance(node, Tag) and in_cover and node.name == "h2":
            if re.match(r"^1\.\s", node.get_text(" ", strip=True)):
                in_cover = False
        if in_cover:
            cover.append(str(node))
        else:
            body.append(str(node))
    return f"<section class='cover'>{''.join(cover)}</section>{''.join(body)}"


def _validate_top_level_nodes(soup: BeautifulSoup) -> None:
    supported = {"h1", "h2", "h3", "p", "ul", "ol", "blockquote", "table"}
    for node in soup.children:
        if isinstance(node, NavigableString):
            if str(node).strip():
                raise ValueError("unsupported top-level Markdown/HTML node: text")
            continue
        if isinstance(node, Tag) and node.name not in supported:
            raise ValueError(f"unsupported top-level Markdown/HTML node: {node.name}")


def build(markdown_path: Path, output_dir: Path) -> tuple[Path, Path]:
    source = markdown_path.read_text(encoding="utf-8")
    body_html = markdown(source, extensions=["tables", "sane_lists"])
    soup = BeautifulSoup(body_html, "html.parser")
    _validate_top_level_nodes(soup)
    output_dir.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.core_properties.title = TITLE
    document.core_properties.subject = SUBJECT
    document.core_properties.author = "MODUA"
    document.core_properties.keywords = "MODUA, 접근성, 생활자립, 상호지원"
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    _configure_page(document)
    _configure_styles(document)
    _append_docx(document, soup)
    _configure_language_metadata(document)
    docx_path = output_dir / f"{OUTPUT_STEM}.docx"
    document.save(docx_path)
    _normalize_docx_archive(docx_path)

    html = (
        "<!doctype html><html lang='ko'><head><meta charset='utf-8'>"
        f"<title>{TITLE}</title><style>{CSS}</style></head>"
        f"<body>{_cover_html(soup)}</body></html>"
    )
    pdf_path = output_dir / f"{OUTPUT_STEM}.pdf"
    identifier = hashlib.sha256(source.encode("utf-8")).hexdigest()
    HTML(string=html, base_url=str(markdown_path.parent)).write_pdf(
        pdf_path,
        pdf_identifier=identifier,
        pdf_tags=True,
    )
    _normalize_pdf(pdf_path, identifier)
    return docx_path, pdf_path


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    outputs = build(
        root / "docs/whitepaper/MODUA_WHITEPAPER_KO.md",
        root / "docs/whitepaper/exports",
    )
    print("\n".join(str(path) for path in outputs))
